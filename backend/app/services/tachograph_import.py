"""
Парсер универсальной CSV/Excel-выгрузки тахографа в записи режима труда и
отдыха (WorkTimeEntry-совместимые данные).

Формат пилотных клиентов на момент написания неизвестен (открытый вопрос
проекта) — реализован обобщённый вариант, наиболее близкий к типичному
экспорту СКЗИ/ЕГТС и стандартным Excel-отчётам по картам водителя:

    ФИО | Номер карты | Тип активности | Начало | Окончание

Колонка "Номер карты" опциональна — сопоставление водителя идёт сначала по
номеру карты (если колонка есть и значение заполнено), затем по ФИО.
Колонка "Окончание" может быть пустой (открытый/текущий интервал).

Формат CSV: разделитель ';', кодировка UTF-8 (с BOM или без) либо CP1251 —
определяется автоматически. Формат XLSX: первый лист, первая строка —
заголовки.

Точка расширения под другие форматы (конкретных пилотных клиентов) —
`_COLUMN_ALIASES` / `_ENTRY_TYPE_ALIASES` и, при необходимости, отдельная
функция `parse_<источник>` рядом с `parse_file`.
"""

import csv
import io
from dataclasses import dataclass
from datetime import datetime

from openpyxl import load_workbook

from app.models.worktime import EntryType

# Заголовок колонки (в нижнем регистре, без пробелов по краям) -> внутреннее имя поля.
_COLUMN_ALIASES: dict[str, str] = {
    "фио": "full_name",
    "водитель": "full_name",
    "фио водителя": "full_name",
    "номер карты": "card_number",
    "карта водителя": "card_number",
    "card number": "card_number",
    "табельный номер": "card_number",
    "тип активности": "entry_type",
    "тип": "entry_type",
    "активность": "entry_type",
    "activity": "entry_type",
    "начало": "start_time",
    "начало периода": "start_time",
    "start": "start_time",
    "окончание": "end_time",
    "конец периода": "end_time",
    "end": "end_time",
}

_ENTRY_TYPE_ALIASES: dict[str, EntryType] = {
    "вождение": EntryType.driving,
    "движение": EntryType.driving,
    "driving": EntryType.driving,
    "отдых": EntryType.rest,
    "перерыв": EntryType.rest,
    "rest": EntryType.rest,
    "break": EntryType.rest,
    "другая работа": EntryType.other_work,
    "работа": EntryType.other_work,
    "other_work": EntryType.other_work,
    "other work": EntryType.other_work,
    "доступность": EntryType.availability,
    "готовность": EntryType.availability,
    "availability": EntryType.availability,
}

_DATETIME_FORMATS = (
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%dT%H:%M:%S",
)

_REQUIRED_FIELDS = ("full_name", "entry_type", "start_time")


@dataclass(frozen=True)
class ParsedRow:
    """Успешно разобранная строка файла — данные готовы к сопоставлению с водителем."""

    row_number: int
    full_name: str
    card_number: str | None
    entry_type: EntryType
    start_time: datetime
    end_time: datetime | None


@dataclass(frozen=True)
class RowError:
    row_number: int
    reason: str


@dataclass(frozen=True)
class ParseResult:
    rows: list[ParsedRow]
    errors: list[RowError]


class UnsupportedFileFormatError(ValueError):
    pass


def parse_file(filename: str, content: bytes) -> ParseResult:
    """Определяет формат по расширению файла и разбирает его в ParseResult."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        return _parse_csv(content)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise UnsupportedFileFormatError(
        f"Неподдерживаемый формат файла: {filename}. Ожидается .csv или .xlsx"
    )


def _parse_csv(content: bytes) -> ParseResult:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251")

    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = list(reader)
    if not raw_rows:
        return ParseResult(rows=[], errors=[])

    header = raw_rows[0]
    field_map = _map_header(header)
    return _build_rows(field_map, raw_rows[1:])


def _parse_xlsx(content: bytes) -> ParseResult:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell) if cell is not None else "" for cell in next(rows_iter)]
    except StopIteration:
        return ParseResult(rows=[], errors=[])

    field_map = _map_header(header)
    raw_rows = [
        ["" if cell is None else cell for cell in row] for row in rows_iter
    ]
    return _build_rows(field_map, raw_rows)


def _map_header(header: list[str]) -> dict[int, str]:
    """Сопоставляет индекс колонки с внутренним именем поля по алиасам заголовка."""
    field_map: dict[int, str] = {}
    for idx, raw_title in enumerate(header):
        title = str(raw_title).strip().lower()
        field = _COLUMN_ALIASES.get(title)
        if field:
            field_map[idx] = field
    return field_map


def _build_rows(field_map: dict[int, str], raw_rows: list[list]) -> ParseResult:
    rows: list[ParsedRow] = []
    errors: list[RowError] = []

    missing_required = [f for f in _REQUIRED_FIELDS if f not in field_map.values()]
    if missing_required:
        errors.append(
            RowError(
                row_number=1,
                reason=(
                    "Не найдены обязательные колонки в заголовке: "
                    + ", ".join(missing_required)
                ),
            )
        )
        return ParseResult(rows=[], errors=errors)

    for offset, raw_row in enumerate(raw_rows):
        row_number = offset + 2  # +1 заголовок, +1 нумерация с 1
        if not raw_row or all(_is_blank(cell) for cell in raw_row):
            continue

        values: dict[str, str] = {}
        for idx, field in field_map.items():
            if idx < len(raw_row):
                values[field] = str(raw_row[idx]).strip() if not _is_blank(raw_row[idx]) else ""

        try:
            parsed = _parse_row(row_number, values)
        except _RowParseError as exc:
            errors.append(RowError(row_number=row_number, reason=str(exc)))
            continue
        rows.append(parsed)

    return ParseResult(rows=rows, errors=errors)


class _RowParseError(ValueError):
    pass


def _parse_row(row_number: int, values: dict[str, str]) -> ParsedRow:
    full_name = values.get("full_name", "")
    if not full_name:
        raise _RowParseError("Не заполнено ФИО водителя")

    entry_type_raw = values.get("entry_type", "")
    entry_type = _resolve_entry_type(entry_type_raw)
    if entry_type is None:
        raise _RowParseError(f"Неизвестный тип активности: '{entry_type_raw}'")

    start_raw = values.get("start_time", "")
    if not start_raw:
        raise _RowParseError("Не заполнено время начала")
    start_time = _parse_datetime(start_raw)
    if start_time is None:
        raise _RowParseError(f"Не удалось разобрать дату/время начала: '{start_raw}'")

    end_raw = values.get("end_time", "")
    end_time = None
    if end_raw:
        end_time = _parse_datetime(end_raw)
        if end_time is None:
            raise _RowParseError(f"Не удалось разобрать дату/время окончания: '{end_raw}'")
        if end_time < start_time:
            raise _RowParseError("Время окончания раньше времени начала")

    card_number = values.get("card_number") or None

    return ParsedRow(
        row_number=row_number,
        full_name=full_name,
        card_number=card_number,
        entry_type=entry_type,
        start_time=start_time,
        end_time=end_time,
    )


def _resolve_entry_type(raw: str) -> EntryType | None:
    key = raw.strip().lower()
    if not key:
        return None
    if key in _ENTRY_TYPE_ALIASES:
        return _ENTRY_TYPE_ALIASES[key]
    try:
        return EntryType(key)
    except ValueError:
        return None


def _parse_datetime(raw: str) -> datetime | None:
    raw = raw.strip()
    if isinstance(raw, datetime):
        return raw
    for fmt in _DATETIME_FORMATS:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return False
